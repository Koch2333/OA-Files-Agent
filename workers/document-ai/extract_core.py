"""Format extractors. Runs inside a child process (see server.py); never imported for parsing by the server."""
import json
import os
import sys

TEXT_CAP = 200000
MAX_XLSX_ROWS = 5000
MAX_PDF_PAGES = 300
MAX_OCR_PAGES = 50
MIN_NATIVE_CHARS = 20
OCR_DPI = 200
HIGH_DPI = 300
MAX_PIXELS = 25_000_000
LOW_CONFIDENCE = 0.6
PROFILES = ("default", "enhanced", "rotate", "deskew", "high_dpi", "binarize")
RENDER_EXTS = {".pdf", ".png", ".jpg", ".jpeg"}
MAX_RENDER_B64 = int(2.5 * 1024 * 1024)
PREVIEW_COLS = 50
PREVIEW_CELL = 200
OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
IMAGE_EXTS = {".png", ".jpg", ".jpeg"}
SUPPORTED_EXTS = {".docx", ".xlsx", ".pdf"} | IMAGE_EXTS


class ExtractError(Exception):
    def __init__(self, code, message=""):
        super().__init__(message or code)
        self.code = code


_ocr_engine = None


def probe():
    """Import-probe only; does not load models."""
    result = {"extraction_available": True, "ocr_available": True}
    try:
        import docx, openpyxl, pymupdf  # noqa: F401
    except Exception:
        result["extraction_available"] = False
    try:
        import rapidocr_onnxruntime  # noqa: F401
    except Exception:
        result["ocr_available"] = False
    return result


def get_ocr():
    global _ocr_engine
    if _ocr_engine is None:
        try:
            from rapidocr_onnxruntime import RapidOCR
            _ocr_engine = RapidOCR()
        except Exception as exc:
            raise ExtractError("OCR_UNAVAILABLE", "OCR engine unavailable") from exc
    return _ocr_engine


def _decode(data):
    import cv2
    import numpy as np
    array = cv2.imdecode(np.frombuffer(data, np.uint8), cv2.IMREAD_COLOR)
    if array is None:
        raise ExtractError("UNREADABLE_FILE", "image cannot be decoded")
    return array


def _engine_lines(array):
    """Run the OCR engine -> list of (y, x, text, confidence)."""
    engine = get_ocr()
    try:
        result, _ = engine(array)
    except Exception as exc:
        raise ExtractError("UNREADABLE_FILE", "OCR failed") from exc
    return [(box[0][1], box[0][0], text, float(conf)) for box, text, conf in (result or [])]


def _assemble(lines):
    """Group boxes into lines by vertical position (result is roughly top-to-bottom) -> (text, mean confidence)."""
    if not lines:
        return "", 0.0
    ordered = sorted(lines, key=lambda t: (round(t[0] / 10), t[1]))
    return "\n".join(t[2] for t in ordered), round(sum(t[3] for t in lines) / len(lines), 4)


def ocr_image_bytes(data):
    """OCR encoded image bytes -> (text, mean confidence). Default profile; behaviour must stay stable."""
    return _assemble(_engine_lines(_decode(data)))


def _cap_pixels(array, limit=MAX_PIXELS):
    import cv2
    h, w = array.shape[:2]
    if h * w <= limit:
        return array
    scale = (limit / float(h * w)) ** 0.5
    return cv2.resize(array, (max(1, int(w * scale)), max(1, int(h * scale))), interpolation=cv2.INTER_AREA)


def _upscale(array, factor=None, min_long=None):
    import cv2
    h, w = array.shape[:2]
    if min_long:
        factor = min_long / float(max(h, w))
    if not factor or factor <= 1.0:
        return array
    factor = min(factor, (MAX_PIXELS / float(h * w)) ** 0.5)  # never allocate beyond the pixel cap
    if factor <= 1.0:
        return array
    return cv2.resize(array, (int(w * factor), int(h * factor)), interpolation=cv2.INTER_CUBIC)


def _enhance(array):
    import cv2
    array = _upscale(array, min_long=2000)
    lab = cv2.cvtColor(array, cv2.COLOR_BGR2LAB)
    l, a, b = cv2.split(lab)
    l = cv2.createCLAHE(clipLimit=2.0, tileGridSize=(8, 8)).apply(l)
    array = cv2.cvtColor(cv2.merge((l, a, b)), cv2.COLOR_LAB2BGR)
    return cv2.fastNlMeansDenoisingColored(array, None, 3, 3, 7, 21)


def _binarize(array):
    import cv2
    gray = cv2.medianBlur(cv2.cvtColor(array, cv2.COLOR_BGR2GRAY), 3)
    binary = cv2.adaptiveThreshold(gray, 255, cv2.ADAPTIVE_THRESH_GAUSSIAN_C, cv2.THRESH_BINARY, 31, 12)
    return cv2.cvtColor(binary, cv2.COLOR_GRAY2BGR)


def _rotate_bound(array, angle):
    """Rotate counter-clockwise by angle degrees on an expanded white canvas."""
    import cv2
    h, w = array.shape[:2]
    m = cv2.getRotationMatrix2D((w / 2.0, h / 2.0), angle, 1.0)
    cos, sin = abs(m[0, 0]), abs(m[0, 1])
    nw, nh = int(h * sin + w * cos), int(h * cos + w * sin)
    m[0, 2] += nw / 2.0 - w / 2.0
    m[1, 2] += nh / 2.0 - h / 2.0
    border = 255 if array.ndim == 2 else (255, 255, 255)
    return cv2.warpAffine(array, m, (nw, nh), flags=cv2.INTER_CUBIC, borderMode=cv2.BORDER_CONSTANT,
                          borderValue=border)


def estimate_skew(array):
    """Projection-profile skew estimate: counter-clockwise degrees that straighten the text, or 0."""
    import cv2
    import numpy as np
    gray = cv2.cvtColor(array, cv2.COLOR_BGR2GRAY)
    h, w = gray.shape
    scale = 700.0 / max(h, w)
    if scale < 1.0:
        gray = cv2.resize(gray, (int(w * scale), int(h * scale)), interpolation=cv2.INTER_AREA)
    _, binary = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY_INV + cv2.THRESH_OTSU)
    if cv2.countNonZero(binary) < 50:
        return 0.0

    def score(angle):
        rotated = _rotate_bound(binary, angle) if angle else binary
        return float(np.var(rotated.sum(axis=1, dtype=np.float64)))

    base = score(0.0)
    best = max(np.arange(-15.0, 15.01, 1.0), key=score)
    best = max(np.arange(best - 1.0, best + 1.01, 0.25), key=score)
    if score(best) < base * 1.03:
        return 0.0
    return float(best)


def _preprocess(array, profile, pdf):
    """Image pipeline for the non-rotate profiles."""
    if profile == "enhanced":
        array = _cap_pixels(_enhance(array))
    elif profile == "deskew":
        angle = estimate_skew(array)
        if 0.5 <= abs(angle) <= 15.0:
            array = _cap_pixels(_rotate_bound(array, angle))
    elif profile == "high_dpi":
        if not pdf:
            array = _upscale(array, factor=2.0)
    elif profile == "binarize":
        array = _binarize(array)
    return array


def _detect_rotation(array):
    """Try 0/90/180/270 (clockwise) on a downscaled copy; best mean confidence * text length wins."""
    import cv2
    h, w = array.shape[:2]
    scale = 900.0 / max(h, w)
    small = cv2.resize(array, (int(w * scale), int(h * scale)), interpolation=cv2.INTER_AREA) if scale < 1.0 else array
    codes = {0: None, 90: cv2.ROTATE_90_CLOCKWISE, 180: cv2.ROTATE_180, 270: cv2.ROTATE_90_COUNTERCLOCKWISE}
    scores = {}
    for angle, code in codes.items():
        view = small if code is None else cv2.rotate(small, code)
        lines = _engine_lines(view)
        chars = sum(len(t[2]) for t in lines)
        scores[angle] = (sum(t[3] for t in lines) / len(lines)) * chars if lines else 0.0
    best = max(scores, key=lambda a: scores[a])
    if best != 0 and scores[best] < scores[0] * 1.25:  # keep the original orientation unless clearly worse
        best = 0
    return best, codes[best]


def ocr_page(data, profile="default", pdf=False):
    """OCR one encoded page image -> dict(text, confidence, lines=[(text, conf)], rotation)."""
    array = _decode(data)
    rotation = 0
    if profile != "default":
        array = _cap_pixels(array)
        if profile == "rotate":
            rotation, code = _detect_rotation(array)
            if code is not None:
                import cv2
                array = cv2.rotate(array, code)
        else:
            array = _preprocess(array, profile, pdf)
    lines = _engine_lines(array)
    text, conf = _assemble(lines)
    return {"text": text, "confidence": conf, "lines": [(t[2], t[3]) for t in lines], "rotation": rotation}


def _render_pdf_page(page, profile):
    """PNG bytes of a PDF page for OCR; pixel count is capped for non-default profiles."""
    dpi = HIGH_DPI if profile == "high_dpi" else OCR_DPI
    if profile != "default":
        rect = page.rect
        pixels = (rect.width * dpi / 72.0) * (rect.height * dpi / 72.0)
        if pixels > MAX_PIXELS:
            dpi = max(36, int(dpi * (MAX_PIXELS / pixels) ** 0.5))
    return page.get_pixmap(dpi=dpi).tobytes("png")


CJK_RANGES = ((0x4E00, 0x9FFF), (0x3400, 0x4DBF), (0xF900, 0xFAFF), (0x20000, 0x2A6DF))
COMMON_PUNCT = set("!\"#$%&'()*+,-./:;<=>?@[\\]^_`{|}~，。、；：？！“”‘’（）【】《》〈〉「」『』…—–·￥％＆＋－／～．")


def _is_cjk(ch):
    code = ord(ch)
    return any(lo <= code <= hi for lo, hi in CJK_RANGES)


def build_quality(records, profile):
    """records: [{page, lines:[(text, conf)], rotation}] for OCR pages. No full text is included."""
    lines = []
    weak = []
    chars = cjk = noise = 0
    rotations = []
    for rec in records:
        page_chars = 0
        for text, conf in rec["lines"]:
            lines.append((rec["page"], text, conf))
            for ch in text:
                if ch.isspace():
                    continue
                chars += 1
                page_chars += 1
                if _is_cjk(ch):
                    cjk += 1
                elif not (ch.isalnum() or ch in COMMON_PUNCT):
                    noise += 1
        confs = [c for _, c in rec["lines"]]
        mean = sum(confs) / len(confs) if confs else 0.0
        if not confs or mean < LOW_CONFIDENCE or page_chars < 5:
            weak.append(rec["page"])
        rotations.append(rec["rotation"])
    low = [t for t in lines if t[2] < LOW_CONFIDENCE]
    worst = sorted(low, key=lambda t: t[2])[:10]
    rotation = 0
    nonzero = [r for r in rotations if r]
    if nonzero:
        rotation = max(set(nonzero), key=nonzero.count)
    return {
        "mean_confidence": round(sum(t[2] for t in lines) / len(lines), 4) if lines else None,
        "low_confidence_ratio": round(len(low) / len(lines), 4) if lines else 0.0,
        "chars": chars,
        "cjk_ratio": round(cjk / chars, 4) if chars else 0.0,
        "noise_ratio": round(noise / chars, 4) if chars else 0.0,
        "lines": len(lines),
        "profile": profile,
        "rotation": rotation,
        "weak_pages": weak,
        "low_confidence_lines": [{"page": p, "text": t[:80], "confidence": round(c, 4)} for p, t, c in worst],
    }


def _meta(title=None, author=None, created=None, modified=None):
    def iso(value):
        return value.isoformat() if hasattr(value, "isoformat") else (value or None)
    return {"title": title or None, "author": author or None, "created": iso(created), "modified": iso(modified)}


def extract_docx(path, ocr):
    import zipfile
    with open(path, "rb") as f:
        head = f.read(8)
    if head == OLE_MAGIC:
        raise ExtractError("ENCRYPTED", "password protected office file")
    if not zipfile.is_zipfile(path):
        raise ExtractError("UNREADABLE_FILE", "not a docx container")
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    try:
        document = docx.Document(path)
        parts = []
        for section in document.sections:
            for header in (section.header, section.first_page_header):
                text = "\n".join(p.text for p in header.paragraphs if p.text.strip())
                for table in header.tables:
                    text += "\n" + _table_text(table)
                if text.strip() and text.strip() not in parts:
                    parts.append(text.strip())
        for child in document.element.body.iterchildren():
            tag = child.tag.rsplit("}", 1)[-1]
            if tag == "p":
                text = Paragraph(child, document).text
                if text.strip():
                    parts.append(text)
            elif tag == "tbl":
                parts.append(_table_text(Table(child, document)))
        for section in document.sections:
            text = "\n".join(p.text for p in section.footer.paragraphs if p.text.strip())
            if text.strip() and text.strip() not in parts:
                parts.append(text.strip())
        core = document.core_properties
        meta = _meta(core.title, core.author, core.created, core.modified)
    except ExtractError:
        raise
    except Exception as exc:
        raise ExtractError("UNREADABLE_FILE", "docx cannot be parsed") from exc
    return {"pages": [{"page": 1, "text": "\n".join(parts), "source": "native", "confidence": 1.0}],
            "ocr_used": False, "warnings": [], "metadata": meta}


def _table_text(table):
    rows = []
    for row in table.rows:
        cells = []
        seen = set()
        for cell in row.cells:
            if id(cell._tc) in seen:
                continue
            seen.add(id(cell._tc))
            cells.append(cell.text.replace("\n", " ").strip())
        if any(cells):
            rows.append("\t".join(cells))
    return "\n".join(rows)


def extract_xlsx(path, ocr):
    import zipfile
    with open(path, "rb") as f:
        head = f.read(8)
    if head == OLE_MAGIC:
        raise ExtractError("ENCRYPTED", "password protected office file")
    if not zipfile.is_zipfile(path):
        raise ExtractError("UNREADABLE_FILE", "not an xlsx container")
    import openpyxl
    warnings = []
    pages = []
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            for index, ws in enumerate(wb.worksheets, 1):
                lines = ["[" + ws.title + "]"]
                count = 0
                for row in ws.iter_rows(values_only=True):
                    count += 1
                    if count > MAX_XLSX_ROWS:
                        warnings.append("ROW_CAP_REACHED:" + ws.title)
                        break
                    cells = ["" if v is None else _cell(v) for v in row]
                    if any(c.strip() for c in cells):
                        lines.append("\t".join(cells).rstrip("\t"))
                pages.append({"page": index, "text": "\n".join(lines), "source": "native", "confidence": 1.0})
            props = wb.properties
            meta = _meta(props.title, props.creator, props.created, props.modified)
        finally:
            wb.close()
    except ExtractError:
        raise
    except Exception as exc:
        raise ExtractError("UNREADABLE_FILE", "xlsx cannot be parsed") from exc
    return {"pages": pages, "ocr_used": False, "warnings": warnings, "metadata": meta}


def _cell(value):
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d") if getattr(value, "hour", 0) == 0 and getattr(value, "minute", 0) == 0 else value.isoformat(sep=" ")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _check_pages(pages, count):
    if pages and (min(pages) < 1 or max(pages) > count):
        raise ExtractError("INVALID_REQUEST", "page out of range")


def extract_pdf(path, ocr, profile="default", pages_filter=None):
    import pymupdf as fitz
    try:
        doc = fitz.open(path)
    except Exception as exc:
        raise ExtractError("UNREADABLE_FILE", "pdf cannot be opened") from exc
    try:
        if doc.needs_pass or doc.is_encrypted:
            raise ExtractError("ENCRYPTED", "pdf is encrypted")
        warnings = []
        pages = []
        ocr_used = False
        total = doc.page_count
        if total == 0:
            raise ExtractError("UNREADABLE_FILE", "pdf has no pages")
        if total > MAX_PDF_PAGES:
            warnings.append("PAGE_CAP_REACHED")
        _check_pages(pages_filter, total)
        records = []
        ocr_done = 0
        for number in range(min(total, MAX_PDF_PAGES)):
            try:
                page = doc[number]
                text = page.get_text("text") or ""
            except Exception as exc:
                raise ExtractError("UNREADABLE_FILE", "pdf page cannot be read") from exc
            if len(text.strip()) >= MIN_NATIVE_CHARS:
                pages.append({"page": number + 1, "text": text, "source": "native", "confidence": 1.0})
                continue
            if pages_filter and (number + 1) not in pages_filter:
                pages.append({"page": number + 1, "text": text, "source": "native", "confidence": 1.0})
                continue
            if not ocr:
                pages.append({"page": number + 1, "text": text, "source": "native", "confidence": 1.0})
                if not text.strip():
                    warnings.append("OCR_DISABLED_EMPTY_PAGE:%d" % (number + 1))
                continue
            if ocr_done >= MAX_OCR_PAGES:
                warnings.append("OCR_PAGE_CAP_REACHED")
                pages.append({"page": number + 1, "text": text, "source": "native", "confidence": 1.0})
                continue
            ocr_done += 1
            try:
                if profile == "default":
                    data = page.get_pixmap(dpi=OCR_DPI).tobytes("png")
                else:
                    data = _render_pdf_page(page, profile)
            except Exception as exc:
                raise ExtractError("UNREADABLE_FILE", "pdf page cannot be rendered") from exc
            res = ocr_page(data, profile, pdf=True)
            ocr_used = True
            records.append({"page": number + 1, "lines": res["lines"], "rotation": res["rotation"]})
            pages.append({"page": number + 1, "text": (text.strip() + "\n" + res["text"]).strip(), "source": "ocr",
                          "confidence": res["confidence"]})
        meta = doc.metadata or {}
        return {"pages": pages, "ocr_used": ocr_used, "warnings": warnings, "_ocr": records,
                "metadata": _meta(meta.get("title"), meta.get("author"), meta.get("creationDate"), meta.get("modDate"))}
    finally:
        doc.close()


def extract_image(path, ocr, profile="default", pages_filter=None):
    _check_pages(pages_filter, 1)
    if not ocr:
        return {"pages": [{"page": 1, "text": "", "source": "native", "confidence": 1.0}], "ocr_used": False,
                "warnings": ["OCR_DISABLED_IMAGE"], "metadata": _meta()}
    with open(path, "rb") as f:
        data = f.read()
    res = ocr_page(data, profile)
    warnings = [] if res["text"].strip() else ["OCR_NO_TEXT"]
    return {"pages": [{"page": 1, "text": res["text"], "source": "ocr", "confidence": res["confidence"]}],
            "ocr_used": True, "warnings": warnings, "metadata": _meta(),
            "_ocr": [{"page": 1, "lines": res["lines"], "rotation": res["rotation"]}]}


def extract(path, ext, ocr, profile="default", pages_filter=None):
    if profile not in PROFILES:
        raise ExtractError("INVALID_REQUEST", "unknown ocr_profile")
    if ext == ".docx":
        result = extract_docx(path, ocr)
    elif ext == ".xlsx":
        result = extract_xlsx(path, ocr)
    elif ext == ".pdf":
        result = extract_pdf(path, ocr, profile, pages_filter)
    elif ext in IMAGE_EXTS:
        result = extract_image(path, ocr, profile, pages_filter)
    else:
        raise ExtractError("UNSUPPORTED_FORMAT", "unsupported format")
    _check_pages(pages_filter, len(result["pages"]))
    records = result.pop("_ocr", None)
    result["quality"] = build_quality(records, profile) if result["ocr_used"] and records else None
    full = "\n\n".join(p["text"] for p in result["pages"] if p["text"])
    result["truncated"] = len(full) > TEXT_CAP
    result["text"] = full[:TEXT_CAP]
    if result["truncated"]:
        for p in result["pages"]:
            p["text"] = p["text"][:TEXT_CAP]
        result["warnings"].append("TEXT_TRUNCATED")
    return result


def _encode_png(img, max_side):
    """PNG-encode a PIL image downscaled to max_side; shrinks dimensions (never switches format) until the base64
    form fits MAX_RENDER_B64."""
    import base64
    import io
    from PIL import Image
    img = img.copy()
    img.thumbnail((max_side, max_side), Image.LANCZOS)
    while True:
        buf = io.BytesIO()
        img.save(buf, format="PNG", optimize=True)
        b64 = base64.b64encode(buf.getvalue())
        if len(b64) <= MAX_RENDER_B64:
            return b64.decode("ascii"), img.width, img.height
        w, h = img.size
        if max(w, h) <= 64:
            raise ExtractError("UNREADABLE_FILE", "page image too large to encode")
        img = img.resize((max(1, int(w * 0.85)), max(1, int(h * 0.85))), Image.LANCZOS)


def render(path, ext, page, max_side):
    """Render one PDF page or an image to PNG -> {image_base64, mime_type, width, height, page, pages}."""
    import io
    from PIL import Image, ImageOps
    if ext not in RENDER_EXTS:
        raise ExtractError("UNSUPPORTED_FORMAT", "format cannot be rendered")
    if ext == ".pdf":
        import pymupdf as fitz
        try:
            doc = fitz.open(path)
        except Exception as exc:
            raise ExtractError("UNREADABLE_FILE", "pdf cannot be opened") from exc
        try:
            if doc.needs_pass or doc.is_encrypted:
                raise ExtractError("ENCRYPTED", "pdf is encrypted")
            pages = doc.page_count
            if pages == 0:
                raise ExtractError("UNREADABLE_FILE", "pdf has no pages")
            if page < 1 or page > pages:
                raise ExtractError("INVALID_REQUEST", "page out of range")
            try:
                pg = doc[page - 1]
                scale = max_side / max(pg.rect.width, pg.rect.height, 1.0)
                data = pg.get_pixmap(matrix=fitz.Matrix(scale, scale)).tobytes("png")
                img = Image.open(io.BytesIO(data))
                img.load()
            except Exception as exc:
                raise ExtractError("UNREADABLE_FILE", "pdf page cannot be rendered") from exc
        finally:
            doc.close()
    else:
        pages = 1
        if page != 1:
            raise ExtractError("INVALID_REQUEST", "page out of range")
        try:
            img = Image.open(path)
            if img.width * img.height > 60_000_000:
                raise ExtractError("UNREADABLE_FILE", "image too large")
            img.load()
            img = ImageOps.exif_transpose(img)
        except ExtractError:
            raise
        except Exception as exc:
            raise ExtractError("UNREADABLE_FILE", "image cannot be decoded") from exc
    if img.mode not in ("RGB", "RGBA", "L"):
        img = img.convert("RGB")
    b64, width, height = _encode_png(img, max_side)
    return {"image_base64": b64, "mime_type": "image/png", "width": width, "height": height, "page": page,
            "pages": pages}


def _clip(text, limit=PREVIEW_CELL):
    text = "" if text is None else str(text)
    return text if len(text) <= limit else text[:limit]


def _heading_level(style):
    import re
    m = re.search(r"(\d+)\s*$", style)
    if style.lower().startswith("title"):
        return 1
    return min(6, max(1, int(m.group(1)))) if m else 1


def preview_docx(path, max_blocks):
    import zipfile
    with open(path, "rb") as f:
        head = f.read(8)
    if head == OLE_MAGIC:
        raise ExtractError("ENCRYPTED", "password protected office file")
    if not zipfile.is_zipfile(path):
        raise ExtractError("UNREADABLE_FILE", "not a docx container")
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    try:
        document = docx.Document(path)
        blocks = []
        truncated = False

        def runs_of(par):
            runs = []
            for r in par.runs:
                if r.text:
                    runs.append({"text": r.text, "bold": bool(r.bold), "italic": bool(r.italic)})
            return runs

        for child in document.element.body.iterchildren():
            tag = child.tag.rsplit("}", 1)[-1]
            if tag not in ("p", "tbl"):
                continue
            if len(blocks) >= max_blocks:
                truncated = True
                break
            if tag == "tbl":
                rows = []
                for row in Table(child, document).rows:
                    cells, seen = [], set()
                    for cell in row.cells:
                        if id(cell._tc) in seen:
                            continue
                        seen.add(id(cell._tc))
                        cells.append(_clip(cell.text.replace("\n", " ").strip()))
                    rows.append(cells)
                if rows:
                    blocks.append({"type": "table", "rows": rows})
                continue
            par = Paragraph(child, document)
            text = par.text
            if not text.strip():
                continue
            style = par.style.name if par.style is not None and par.style.name else ""
            if style.lower().startswith(("heading", "title")) or style.startswith("标题"):
                blocks.append({"type": "heading", "level": _heading_level(style), "text": text.strip()})
                continue
            ppr = child.pPr
            num = ppr.numPr if ppr is not None else None
            if "list" in style.lower() or "列表" in style or num is not None:
                level = 0
                if num is not None and num.ilvl is not None:
                    level = int(num.ilvl.val)
                blocks.append({"type": "list_item", "level": level, "runs": runs_of(par)})
                continue
            align = "left"
            if par.alignment == WD_ALIGN_PARAGRAPH.CENTER:
                align = "center"
            elif par.alignment == WD_ALIGN_PARAGRAPH.RIGHT:
                align = "right"
            blocks.append({"type": "paragraph", "runs": runs_of(par), "align": align})
        header = footer = ""
        for section in document.sections:
            h = "\n".join(p.text for p in section.header.paragraphs if p.text.strip())
            f_ = "\n".join(p.text for p in section.footer.paragraphs if p.text.strip())
            header = header or h.strip()
            footer = footer or f_.strip()
        images = sum(1 for el in document.element.body.iter() if el.tag.rsplit("}", 1)[-1] == "blip")
        warnings = ["含 %d 张图片未显示" % images] if images else []
    except ExtractError:
        raise
    except Exception as exc:
        raise ExtractError("UNREADABLE_FILE", "docx cannot be parsed") from exc
    return {"kind": "document", "blocks": blocks, "header": header, "footer": footer, "truncated": truncated,
            "warnings": warnings}


def preview_xlsx(path, max_rows):
    import zipfile
    from itertools import islice, zip_longest
    with open(path, "rb") as f:
        head = f.read(8)
    if head == OLE_MAGIC:
        raise ExtractError("ENCRYPTED", "password protected office file")
    if not zipfile.is_zipfile(path):
        raise ExtractError("UNREADABLE_FILE", "not an xlsx container")
    import openpyxl
    sheets = []
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        wf = openpyxl.load_workbook(path, read_only=True, data_only=False)
        try:
            for ws in wb.worksheets:
                if getattr(ws, "sheet_state", "visible") != "visible":
                    continue
                wsf = wf[ws.title]
                rows = []
                for vals, forms in islice(zip_longest(ws.iter_rows(values_only=True), wsf.iter_rows(values_only=True)),
                                          max_rows):
                    vals = vals or ()
                    forms = forms or ()
                    n = max(len(vals), len(forms))
                    row = []
                    for i in range(min(n, PREVIEW_COLS)):
                        v = vals[i] if i < len(vals) else None
                        if v is None and i < len(forms) and isinstance(forms[i], str) and forms[i].startswith("="):
                            v = forms[i]
                        row.append("" if v is None else _clip(_cell(v)))
                    rows.append(row)
                rows_total = ws.max_row or len(rows)
                cols_total = ws.max_column or max((len(r) for r in rows), default=0)
                rows_total = max(rows_total, len(rows))
                sheets.append({"name": ws.title, "rows": rows, "rows_total": rows_total, "cols_total": cols_total,
                               "truncated": rows_total > max_rows or cols_total > PREVIEW_COLS})
        finally:
            wb.close()
            wf.close()
    except ExtractError:
        raise
    except Exception as exc:
        raise ExtractError("UNREADABLE_FILE", "xlsx cannot be parsed") from exc
    return {"kind": "spreadsheet", "sheets": sheets, "truncated": any(s["truncated"] for s in sheets),
            "warnings": []}


def preview(path, ext, max_rows=200, max_blocks=600):
    if ext == ".docx":
        return preview_docx(path, max_blocks)
    if ext == ".xlsx":
        return preview_xlsx(path, max_rows)
    if ext == ".pdf":
        import pymupdf as fitz
        try:
            doc = fitz.open(path)
        except Exception as exc:
            raise ExtractError("UNREADABLE_FILE", "pdf cannot be opened") from exc
        try:
            if doc.needs_pass or doc.is_encrypted:
                raise ExtractError("ENCRYPTED", "pdf is encrypted")
            if doc.page_count == 0:
                raise ExtractError("UNREADABLE_FILE", "pdf has no pages")
            return {"kind": "pdf", "pages": doc.page_count}
        finally:
            doc.close()
    if ext in IMAGE_EXTS:
        return {"kind": "image", "pages": 1}
    raise ExtractError("UNSUPPORTED_FORMAT", "unsupported format")


def main():
    # argv: path ext ocr(0/1) [profile] [pages csv] | --render path ext page max_side | --preview path ext max_rows
    # max_blocks | --probe. Result JSON goes to stdout as one UTF-8 line.
    out = sys.stdout.buffer
    if sys.argv[1] == "--probe":
        out.write(json.dumps(probe()).encode("utf-8") + b"\n")
        out.flush()
        return
    try:
        if sys.argv[1] == "--render":
            body = {"data": render(sys.argv[2], sys.argv[3], int(sys.argv[4]), int(sys.argv[5]))}
        elif sys.argv[1] == "--preview":
            body = {"data": preview(sys.argv[2], sys.argv[3], int(sys.argv[4]), int(sys.argv[5]))}
        else:
            path, ext, ocr = sys.argv[1], sys.argv[2], sys.argv[3] == "1"
            profile = sys.argv[4] if len(sys.argv) > 4 else "default"
            pages_filter = [int(x) for x in sys.argv[5].split(",") if x] if len(sys.argv) > 5 else None
            body = {"data": extract(path, ext, ocr, profile, pages_filter)}
    except ExtractError as exc:
        body = {"error": {"code": exc.code, "message": str(exc)}}
    except MemoryError:
        body = {"error": {"code": "UNREADABLE_FILE", "message": "out of memory"}}
    except Exception:
        body = {"error": {"code": "UNREADABLE_FILE", "message": "extraction failed"}}
    out.write(json.dumps(body, ensure_ascii=False).encode("utf-8") + b"\n")
    out.flush()


if __name__ == "__main__":
    main()
