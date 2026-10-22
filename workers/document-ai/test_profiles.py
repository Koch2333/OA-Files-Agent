"""Tests for OCR profiles, quality metrics, /v1/render, /v1/preview and extra allowed roots."""
import base64
import io
import json
import os
import shutil
import struct
import sys
import tempfile
import threading
import unittest

import server
from server import create_server
from test_server import SAMPLES, TOKEN, ServerCase

DEGRADED = os.path.join(SAMPLES, "degraded")
QUALITY_KEYS = {"mean_confidence", "low_confidence_ratio", "chars", "cjk_ratio", "noise_ratio", "lines", "profile",
                "rotation", "weak_pages", "low_confidence_lines"}


def png_size(b64):
    raw = base64.b64decode(b64)
    assert raw[:8] == b"\x89PNG\r\n\x1a\n"
    return struct.unpack(">II", raw[16:24])


class ChildCommandTest(unittest.TestCase):
    def test_frozen_modes(self):
        self.assertEqual(server.child_command("extract", frozen=True, executable="w.exe"), ["w.exe", "--extract-child"])
        self.assertEqual(server.child_command("render", frozen=True, executable="w.exe"), ["w.exe", "--render-child"])
        self.assertEqual(server.child_command("preview", frozen=True, executable="w.exe"),
                         ["w.exe", "--preview-child"])

    def test_source_modes(self):
        core = server.CORE
        self.assertEqual(server.child_command("extract", frozen=False, executable="py"), ["py", "-I", core])
        self.assertEqual(server.child_command("render", frozen=False, executable="py"), ["py", "-I", core, "--render"])
        self.assertEqual(server.child_command("preview", frozen=False, executable="py"),
                         ["py", "-I", core, "--preview"])

    def test_unknown_mode(self):
        with self.assertRaises(ValueError):
            server.child_command("bogus")

    def test_worker_main_maps_child_flags(self):
        import worker_main
        self.assertEqual(worker_main.CHILD_MODES["--render-child"], "--render")
        self.assertEqual(worker_main.CHILD_MODES["--preview-child"], "--preview")
        self.assertIsNone(worker_main.CHILD_MODES["--extract-child"])


class ProfileRequestTest(ServerCase):
    def req(self, name, **extra):
        body = {"path": self.sample(name), "filename": name, "ocr": True}
        body.update(extra)
        return self.call("POST", "/v1/extract", body)

    def test_invalid_fields(self):
        for extra in ({"ocr_profile": "nope"}, {"ocr_profile": 3}, {"pages": "1"}, {"pages": [0]}, {"pages": ["1"]},
                      {"pages": [True]}, {"pages": [1.5]}, {"pages": list(range(1, 52))}):
            status, body = self.req("办公电脑采购合同.pdf", **extra)
            self.assertEqual((status, body["error"]["code"]), (400, "INVALID_REQUEST"), extra)

    def test_page_out_of_range(self):
        status, body = self.req("办公电脑采购合同.pdf", pages=[5])
        self.assertEqual((status, body["error"]["code"]), (400, "INVALID_REQUEST"))

    def test_native_pdf_has_null_quality(self):
        status, body = self.req("办公电脑采购合同.pdf", ocr_profile="enhanced", pages=[1])
        self.assertEqual(status, 200, body)
        self.assertIsNone(body["data"]["quality"])

    def test_docx_has_null_quality(self):
        status, body = self.req("财务部_2026年3月报销申请.docx")
        self.assertEqual(status, 200)
        self.assertIsNone(body["data"]["quality"])

    def test_quality_for_ocr_image(self):
        status, body = self.req("采购申请单_拍照.png")
        self.assertEqual(status, 200, body)
        quality = body["data"]["quality"]
        self.assertEqual(set(quality), QUALITY_KEYS)
        self.assertEqual(quality["profile"], "default")
        self.assertEqual(quality["rotation"], 0)
        self.assertGreater(quality["chars"], 10)
        self.assertGreater(quality["cjk_ratio"], 0.3)
        self.assertGreater(quality["mean_confidence"], 0.5)
        self.assertLessEqual(len(quality["low_confidence_lines"]), 10)
        self.assertNotIn("采购申请单", json.dumps(quality, ensure_ascii=False)
                         if quality["low_confidence_lines"] == [] else "")

    def test_quality_for_scanned_pdf_with_profile_and_pages(self):
        status, body = self.req("劳动合同续签_扫描件.pdf", ocr_profile="enhanced", pages=[1])
        self.assertEqual(status, 200, body)
        self.assertEqual(body["data"]["quality"]["profile"], "enhanced")
        self.assertIn("劳动合同续签协议", body["data"]["text"])

    def test_pages_limit_ocr(self):
        # a one-page scan with pages=[1] OCRs; (out of range checked above)
        status, body = self.req("scan_20260520_001.pdf", ocr_profile="deskew", pages=[1])
        self.assertEqual(status, 200, body)
        self.assertTrue(body["data"]["ocr_used"])

    def test_other_profiles_accepted(self):
        for profile in ("high_dpi", "binarize"):
            status, body = self.req("人事部_入职培训通知.jpg", ocr_profile=profile)
            self.assertEqual(status, 200, (profile, body))
            self.assertEqual(body["data"]["quality"]["profile"], profile)


class QualityUnitTest(unittest.TestCase):
    def test_metrics(self):
        import extract_core
        records = [{"page": 1, "rotation": 90, "lines": [("合同ABC123", 0.9), ("☆☆♪♪※※", 0.3), ("y" * 100, 0.5), ("z" * 10, 0.95)]},
                   {"page": 2, "rotation": 0, "lines": [("ab", 0.95)]}]
        q = extract_core.build_quality(records, "rotate")
        self.assertEqual(q["lines"], 5)
        self.assertEqual(q["rotation"], 90)
        self.assertEqual(q["weak_pages"], [2])
        self.assertAlmostEqual(q["low_confidence_ratio"], 0.4)
        self.assertGreater(q["noise_ratio"], 0)
        self.assertEqual(len(q["low_confidence_lines"][1]["text"]), 80)
        self.assertEqual(q["low_confidence_lines"][0]["confidence"], 0.3)
        empty = extract_core.build_quality([{"page": 1, "rotation": 0, "lines": []}], "default")
        self.assertIsNone(empty["mean_confidence"])
        self.assertEqual(empty["weak_pages"], [1])


class RotationRescueTest(unittest.TestCase):
    """Real OCR on this machine. Measured finding (see README): rapidocr already reads 90/180 degree pages with the
    default profile (recall 1.0), so `rotate` cannot beat it; it must at least never be worse and must recover the
    keywords. If a future engine/model loses that ability, `rotate` is the profile expected to recover it."""

    def recall(self, name, profile, keywords):
        import extract_core
        path = os.path.join(DEGRADED, name)
        text = extract_core.extract(path, ".png", True, profile)["text"]
        flat = "".join(text.split())
        return sum(1 for k in keywords if k in flat) / float(len(keywords))

    def test_rotate_recovers_rotated_pages(self):
        import extract_core
        if not extract_core.probe()["ocr_available"]:
            self.skipTest("OCR engine (rapidocr_onnxruntime) not installed")
        if not os.path.isdir(DEGRADED):
            self.skipTest("run make_degraded.py to create tests/samples/degraded")
        with open(os.path.join(DEGRADED, "manifest.json"), encoding="utf-8") as f:
            entries = {e["file"]: e for e in json.load(f)["files"]}
        for name in ("purchase_photo_rot180.png", "purchase_photo_rot90.png"):
            keywords = entries[name]["keywords"]
            base = self.recall(name, "default", keywords)
            rescued = self.recall(name, "rotate", keywords)
            self.assertGreaterEqual(rescued, base, "%s: rotate %.2f vs default %.2f" % (name, rescued, base))
            self.assertGreaterEqual(rescued, 0.5, "%s: rotate recall %.2f" % (name, rescued))


class RenderTest(ServerCase):
    def render(self, name, **extra):
        body = {"path": self.sample(name)}
        body.update(extra)
        return self.call("POST", "/v1/render", body)

    def test_pdf(self):
        status, body = self.render("办公电脑采购合同.pdf", max_side=600)
        self.assertEqual(status, 200, body)
        data = body["data"]
        w, h = png_size(data["image_base64"])
        self.assertEqual((w, h), (data["width"], data["height"]))
        self.assertLessEqual(max(w, h), 600)
        self.assertGreater(max(w, h), 500)
        self.assertEqual((data["mime_type"], data["page"], data["pages"]), ("image/png", 1, 1))

    def test_image_not_upscaled(self):
        status, body = self.render("采购申请单_拍照.png", max_side=2000)
        self.assertEqual(status, 200, body)
        w, h = png_size(body["data"]["image_base64"])
        self.assertLessEqual(max(w, h), 2000)
        self.assertLessEqual(len(body["data"]["image_base64"]), int(2.5 * 1024 * 1024))

    def test_default_max_side(self):
        status, body = self.render("办公电脑采购合同.pdf")
        self.assertEqual(status, 200)
        self.assertLessEqual(max(body["data"]["width"], body["data"]["height"]), 1400)

    def test_errors(self):
        status, body = self.render("财务部_2026年3月报销申请.docx")
        self.assertEqual((status, body["error"]["code"]), (422, "UNSUPPORTED_FORMAT"))
        status, body = self.render("办公电脑采购合同.pdf", page=2)
        self.assertEqual((status, body["error"]["code"]), (400, "INVALID_REQUEST"))
        status, body = self.render("采购申请单_拍照.png", page=2)
        self.assertEqual((status, body["error"]["code"]), (400, "INVALID_REQUEST"))
        for extra in ({"page": 0}, {"page": "1"}, {"max_side": 100}, {"max_side": 5000}, {"max_side": True}):
            status, body = self.render("办公电脑采购合同.pdf", **extra)
            self.assertEqual((status, body["error"]["code"]), (400, "INVALID_REQUEST"), extra)
        status, body = self.render("加密合同.pdf")
        self.assertEqual((status, body["error"]["code"]), (422, "ENCRYPTED"))
        status, body = self.render("不存在.pdf")
        self.assertEqual((status, body["error"]["code"]), (404, "FILE_NOT_FOUND"))

    def test_requires_auth(self):
        self.assertEqual(self.call("POST", "/v1/render", {"path": "x"}, auth=False)[0], 401)
        self.assertEqual(self.call("POST", "/v1/preview", {"path": "x"}, auth=False)[0], 401)


class RootsBase(ServerCase):
    spool = None
    more_roots = ()

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp()
        cls.spool = os.path.join(cls.tmp, "spool")
        cls.extra = os.path.join(cls.tmp, "archive")
        cls.outside = os.path.join(cls.tmp, "outside")
        for d in (cls.spool, cls.extra, cls.outside):
            os.makedirs(d)
        cls.server = create_server(0, TOKEN, spool_root=cls.spool, timeout=cls.timeout, extra_roots=[cls.extra, *cls.more_roots])
        cls.thread = threading.Thread(target=cls.server.serve_forever, daemon=True)
        cls.thread.start()
        cls.base = "http://127.0.0.1:%d" % cls.server.server_port

    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def preview(self, path, **extra):
        body = {"path": path}
        body.update(extra)
        return self.call("POST", "/v1/preview", body)


class AllowedRootsTest(RootsBase):
    def test_roots(self):
        pdf = self.sample("办公电脑采购合同.pdf")
        for folder in (self.spool, self.extra, self.outside):
            shutil.copy(pdf, os.path.join(folder, "a.pdf"))
        self.assertEqual(self.preview(os.path.join(self.spool, "a.pdf"))[0], 200)
        status, body = self.preview(os.path.join(self.extra, "a.pdf"))
        self.assertEqual((status, body["data"]), (200, {"kind": "pdf", "pages": 1}))
        status, body = self.preview(os.path.join(self.outside, "a.pdf"))
        self.assertEqual((status, body["error"]["code"]), (403, "PATH_NOT_ALLOWED"))
        status, body = self.preview(os.path.join(self.extra, "..", "outside", "a.pdf"))
        self.assertEqual((status, body["error"]["code"]), (403, "PATH_NOT_ALLOWED"))
        # render also honours the extra root; extract stays spool-only
        status, _ = self.call("POST", "/v1/render", {"path": os.path.join(self.extra, "a.pdf"), "max_side": 300})
        self.assertEqual(status, 200)
        status, body = self.extract(os.path.join(self.extra, "a.pdf"))
        self.assertEqual((status, body["error"]["code"]), (403, "PATH_NOT_ALLOWED"))
        status, _ = self.extract(os.path.join(self.spool, "a.pdf"))
        self.assertEqual(status, 200)

    def test_symlink_escape_from_extra_root(self):
        target = os.path.join(self.outside, "s.pdf")
        shutil.copy(self.sample("办公电脑采购合同.pdf"), target)
        try:
            os.symlink(target, os.path.join(self.extra, "link.pdf"))
        except (OSError, NotImplementedError):
            self.skipTest("symlinks not permitted on this system")
        status, body = self.preview(os.path.join(self.extra, "link.pdf"))
        self.assertEqual((status, body["error"]["code"]), (403, "PATH_NOT_ALLOWED"))

    def test_parse_roots(self):
        raw = os.pathsep.join([self.extra, "relative", "", self.outside])
        self.assertEqual(server.parse_roots(raw), [self.extra, self.outside])


class PreviewTest(RootsBase):
    more_roots = (SAMPLES,)

    def test_docx_blocks(self):
        import docx
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        d = docx.Document()
        d.add_heading("总标题", level=1)
        d.add_heading("小节", level=2)
        p = d.add_paragraph()
        p.add_run("粗体").bold = True
        p.add_run(" 普通 ")
        p.add_run("斜体").italic = True
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        d.add_paragraph("项目一", style="List Bullet")
        t = d.add_table(rows=2, cols=2)
        for i, row in enumerate(t.rows):
            for j, cell in enumerate(row.cells):
                cell.text = "r%dc%d" % (i, j)
        d.sections[0].header.paragraphs[0].text = "页眉文字"
        d.sections[0].footer.paragraphs[0].text = "页脚文字"
        path = os.path.join(self.extra, "p.docx")
        d.save(path)
        status, body = self.preview(path)
        self.assertEqual(status, 200, body)
        data = body["data"]
        self.assertEqual(data["kind"], "document")
        types = [b["type"] for b in data["blocks"]]
        self.assertEqual(types, ["heading", "heading", "paragraph", "list_item", "table"])
        self.assertEqual([b["level"] for b in data["blocks"][:2]], [1, 2])
        para = data["blocks"][2]
        self.assertEqual(para["align"], "center")
        self.assertEqual([(r["text"], r["bold"], r["italic"]) for r in para["runs"]],
                         [("粗体", True, False), (" 普通 ", False, False), ("斜体", False, True)])
        self.assertEqual(data["blocks"][4]["rows"], [["r0c0", "r0c1"], ["r1c0", "r1c1"]])
        self.assertEqual((data["header"], data["footer"]), ("页眉文字", "页脚文字"))
        self.assertIs(data["truncated"], False)

    def test_docx_truncation_by_max_blocks(self):
        status, body = self.preview(self.sample("财务部_2026年3月报销申请.docx"), max_blocks=1)
        self.assertEqual(status, 200)
        self.assertTrue(body["data"]["truncated"])
        self.assertEqual(len(body["data"]["blocks"]), 1)

    def test_xlsx_truncation_and_hidden_sheet(self):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "数据"
        for i in range(1, 31):
            ws.append([i, "x" * 300, "=A%d*2" % i] + list(range(60)))
        hidden = wb.create_sheet("隐藏")
        hidden["A1"] = "secret"
        hidden.sheet_state = "hidden"
        path = os.path.join(self.extra, "p.xlsx")
        wb.save(path)
        status, body = self.preview(path, max_rows=10)
        self.assertEqual(status, 200, body)
        data = body["data"]
        self.assertEqual(data["kind"], "spreadsheet")
        self.assertEqual([s["name"] for s in data["sheets"]], ["数据"])
        sheet = data["sheets"][0]
        self.assertEqual(len(sheet["rows"]), 10)
        self.assertEqual(sheet["rows_total"], 30)
        self.assertEqual(sheet["cols_total"], 63)
        self.assertTrue(sheet["truncated"] and data["truncated"])
        self.assertTrue(all(len(r) <= 50 for r in sheet["rows"]))
        self.assertTrue(all(len(c) <= 200 for r in sheet["rows"] for c in r))
        self.assertEqual(sheet["rows"][0][2], "=A1*2")  # no cached value -> formula text

    def test_xlsx_sample(self):
        status, body = self.preview(self.sample("采购部_2026年Q1设备采购报表.xlsx"))
        self.assertEqual(status, 200, body)
        self.assertEqual(len(body["data"]["sheets"]), 2)
        self.assertFalse(body["data"]["truncated"])

    def test_pdf_and_image_kinds(self):
        status, body = self.preview(self.sample("办公电脑采购合同.pdf"))
        self.assertEqual((status, body["data"]), (200, {"kind": "pdf", "pages": 1}))
        status, body = self.preview(self.sample("采购申请单_拍照.png"))
        self.assertEqual((status, body["data"]), (200, {"kind": "image", "pages": 1}))

    def test_errors(self):
        legacy = os.path.join(self.extra, "old.doc")
        with open(legacy, "wb") as f:
            f.write(b"x")
        status, body = self.preview(legacy)
        self.assertEqual((status, body["error"]["code"]), (422, "UNSUPPORTED_FORMAT"))
        status, body = self.preview(self.sample("财务部_损坏的报表.docx"))
        self.assertEqual((status, body["error"]["code"]), (422, "UNREADABLE_FILE"))
        status, body = self.preview(self.sample("加密合同.pdf"))
        self.assertEqual((status, body["error"]["code"]), (422, "ENCRYPTED"))
        status, body = self.preview(self.sample("办公电脑采购合同.pdf"), max_rows=0)
        self.assertEqual((status, body["error"]["code"]), (400, "INVALID_REQUEST"))
