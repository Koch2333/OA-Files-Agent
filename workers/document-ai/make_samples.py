"""Generate Chinese OA sample fixtures into tests/samples/ plus manifest.json (ground truth).

Run with the worker venv:  .venv\\Scripts\\python.exe make_samples.py
Needs a CJK font (default C:\\Windows\\Fonts\\msyh.ttc; override with OA_CJK_FONT).
"""
import datetime
import io
import json
import os

import docx
import pymupdf as fitz
import openpyxl
from docx.shared import Pt
from PIL import Image, ImageDraw, ImageFont

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.normpath(os.path.join(HERE, "..", "..", "tests", "samples"))
FONT = os.environ.get("OA_CJK_FONT", r"C:\Windows\Fonts\msyh.ttc")
MANIFEST = []


def record(name, department, doc_type, date, **extra):
    entry = {"file": name, "department": department, "type": doc_type, "date": date}
    entry.update(extra)
    MANIFEST.append(entry)


def path(name):
    return os.path.join(OUT, name)


def make_docx(name, lines, table=None, header=None):
    d = docx.Document()
    d.styles["Normal"].font.name = "Microsoft YaHei"
    d.styles["Normal"].font.size = Pt(11)
    if header:
        d.sections[0].header.paragraphs[0].text = header
    for i, line in enumerate(lines):
        d.add_heading(line, level=1) if i == 0 else d.add_paragraph(line)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        t.style = "Table Grid"
        for r, row in enumerate(table):
            for c, val in enumerate(row):
                t.cell(r, c).text = val
    d.core_properties.author = "OA系统"
    d.save(path(name))


def make_xlsx(name, sheets):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets:
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(row)
    wb.save(path(name))


def render_page(lines, size=(1240, 1754), font_size=34):
    img = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(img)
    title_font = ImageFont.truetype(FONT, int(font_size * 1.5))
    body_font = ImageFont.truetype(FONT, font_size)
    y = 120
    for i, line in enumerate(lines):
        draw.text((90, y), line, fill="black", font=title_font if i == 0 else body_font)
        y += int(font_size * (2.4 if i == 0 else 1.8))
    return img


def make_native_pdf(name, lines, encrypt=False):
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_font(fontname="cjk", fontfile=FONT)
    y = 90
    for i, line in enumerate(lines):
        page.insert_text((60, y), line, fontname="cjk", fontsize=18 if i == 0 else 12)
        y += 40 if i == 0 else 26
    if encrypt:
        pdf.save(path(name), encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw="owner-secret", user_pw="user-secret")
    else:
        pdf.save(path(name))


def make_scanned_pdf(name, lines):
    img = render_page(lines)
    buf = io.BytesIO()
    img.save(buf, "PNG")
    pdf = fitz.open()
    page = pdf.new_page(width=595, height=842)
    page.insert_image(page.rect, stream=buf.getvalue())
    pdf.save(path(name))


def make_image(name, lines, fmt):
    img = render_page(lines, size=(1000, 800), font_size=36)
    img.save(path(name), fmt, **({"quality": 90} if fmt == "JPEG" else {}))


def main():
    os.makedirs(OUT, exist_ok=True)

    name = "财务部_2026年3月报销申请.docx"
    make_docx(name, ["费用报销申请单", "申请部门：财务部", "申请人：李明", "申请日期：2026年3月15日",
                     "报销事由：2026年3月差旅及办公费用报销，请予审批。"],
              table=[["序号", "项目", "金额(元)"], ["1", "差旅费", "3280.00"], ["2", "办公用品", "560.50"]])
    record(name, "财务部", "申请", "2026-03-15", keywords=["费用报销申请单", "3280.00"])

    name = "关于2026年国庆节放假安排的通知.docx"
    make_docx(name, ["关于2026年国庆节放假安排的通知", "发文部门：行政部", "发文日期：2026-09-28",
                     "各部门：根据国家规定，现将国庆节放假安排通知如下。"],
              table=[["日期", "安排"], ["10月1日至10月7日", "放假调休"], ["10月8日", "正常上班"]],
              header="行政部文件 · 内部发布")
    record(name, "行政部", "通知", "2026-09-28", keywords=["国庆节放假安排", "放假调休", "行政部文件"])

    name = "采购部_2026年Q1设备采购报表.xlsx"
    make_xlsx(name, [("采购明细", [["设备名称", "数量", "单价", "供应商"], ["笔记本电脑", 20, 6800, "联想科技"],
                                  ["显示器", 30, 1200, "戴尔中国"]]),
                     ("汇总", [["编制日期", "2026-03-31"], ["合计金额", 172000]])])
    record(name, "采购部", "报表", "2026-03-31", keywords=["采购明细", "联想科技", "汇总"])

    name = "2026年2月考勤汇总报表.xlsx"
    make_xlsx(name, [("考勤", [["编制部门：人事部"], ["统计日期：2026年2月28日"], ["姓名", "出勤天数", "迟到次数"],
                              ["王芳", 20, 0], ["赵强", 19, 2]])])
    record(name, "人事部", "报表", "2026-02-28", keywords=["考勤", "王芳", "迟到次数"])

    name = "办公电脑采购合同.pdf"
    make_native_pdf(name, ["办公电脑采购合同", "合同编号：CG-2026-0315", "甲方：本公司（经办部门：采购部）",
                           "乙方：联想科技有限公司", "签订日期：2026年3月15日",
                           "第一条 乙方向甲方供应笔记本电脑20台，总价136000元。"])
    record(name, "采购部", "合同", "2026-03-15", keywords=["办公电脑采购合同", "联想科技有限公司"])

    name = "行政部_会议室使用申请.pdf"
    make_native_pdf(name, ["会议室使用申请", "申请人：陈静", "使用时间：2026-04-02 14:00-16:00",
                           "事由：季度工作例会，预计参会15人。"])
    record(name, "行政部", "申请", "2026-04-02", keywords=["会议室使用申请", "季度工作例会"])

    name = "劳动合同续签_扫描件.pdf"
    make_scanned_pdf(name, ["劳动合同续签协议", "用人部门：人事部", "员工姓名：张伟",
                            "签订日期：2026年1月5日", "合同期限：三年，自2026年2月1日起算。"])
    record(name, "人事部", "合同", "2026-01-05", keywords=["劳动合同续签协议", "张伟"], scanned=True)

    name = "scan_20260520_001.pdf"
    make_scanned_pdf(name, ["关于调整报销流程的通知", "发文部门：财务部", "发文日期：2026-05-20",
                            "各部门：自6月1日起，报销单须经部门负责人审批。"])
    record(name, "财务部", "通知", "2026-05-20", keywords=["调整报销流程", "部门负责人"], scanned=True)

    name = "采购申请单_拍照.png"
    make_image(name, ["采购申请单", "申请部门：采购部", "申请日期：2026年6月8日", "物品：打印机 5台"], "PNG")
    record(name, "采购部", "申请", "2026-06-08", keywords=["采购申请单", "打印机"], scanned=True)

    name = "人事部_入职培训通知.jpg"
    make_image(name, ["入职培训通知", "新员工请于2026-07-01上午9点", "到三楼培训室参加入职培训。"], "JPEG")
    record(name, "人事部", "通知", "2026-07-01", keywords=["入职培训通知", "培训室"], scanned=True)

    name = "财务部_损坏的报表.docx"
    with open(path(name), "wb") as f:
        f.write(b"PK\x03\x04 this is not a real docx" + os.urandom(64))
    record(name, "财务部", "报表", None, expected_error="UNREADABLE_FILE")

    name = "加密合同.pdf"
    make_native_pdf(name, ["保密合同", "甲方：本公司", "签订日期：2026年8月1日"], encrypt=True)
    record(name, None, "合同", "2026-08-01", expected_error="ENCRYPTED")

    with open(path("manifest.json"), "w", encoding="utf-8") as f:
        json.dump({"generated": datetime.date.today().isoformat(), "files": MANIFEST}, f, ensure_ascii=False, indent=2)
    print("wrote", len(MANIFEST), "samples to", OUT)


if __name__ == "__main__":
    main()
